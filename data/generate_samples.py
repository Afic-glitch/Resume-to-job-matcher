import os
from pathlib import Path
from docx import Document

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "sample_data"
DATA_DIR.mkdir(parents=True, exist_ok=True)

# 1. Create Sample DOCX Resume (Priya Sharma - CAND-027)
doc = Document()
doc.add_heading('Priya Sharma', level=0)
doc.add_paragraph('priya.sharma@example.com | +91 9811122233 | Bengaluru, India')
doc.add_paragraph('Female | Date of Birth: 15/08/2001')

doc.add_heading('Professional Summary', level=1)
doc.add_paragraph(
    'Backend Engineer with 1.5 years of industry experience specializing in Python, '
    'FastAPI, and SQL database performance optimization. Built scalable RESTful microservices '
    'and implemented automated CI/CD workflows using Git.'
)

doc.add_heading('Education', level=1)
doc.add_paragraph('B.Tech in Computer Science & Engineering - Batch 2024')

doc.add_heading('Technical Skills', level=1)
doc.add_paragraph('Python, SQL, REST API, Git, PostgreSQL, Django, Linux, HTML, CSS')

doc.add_heading('Work Experience', level=1)
doc.add_paragraph(
    'Software Engineer at TechSprint Solutions (1.5 years):\n'
    '- Developed robust REST APIs using FastAPI and Python.\n'
    '- Managed PostgreSQL relational schemas and resolved query latency bottlenecks.\n'
    '- Collaborated with team members using Git version control.'
)

doc.add_heading('Projects', level=1)
doc.add_paragraph(
    '1. High-Throughput Checkout API: Built asynchronous REST endpoints handling 1,500 req/sec.\n'
    '2. Database Migration Engine: Automated schema versioning and data transfer.'
)

doc.add_heading('Certifications', level=1)
doc.add_paragraph('Python Professional Developer Certificate')

docx_path = DATA_DIR / "sample_resume_priya_sharma.docx"
doc.save(str(docx_path))
print(f"Created DOCX: {docx_path}")

# 2. Create Sample TXT files for Arun Verma and Rohan Mehta
c1_txt = """
Arun Verma
arun.verma@example.com | +91 9876543210 | Male
New Delhi, India

Summary:
Full-stack Python Backend Engineer with 2 years of experience developing microservices.
Skills: Python, SQL, REST API, Git, Docker, AWS, PostgreSQL, Linux
Education: B.Tech Computer Science, Tier-1 Institute
Experience: 2 years as Junior Software Engineer at CloudCorp. Built RESTful microservices using FastAPI and SQLAlchemy. Deployed containers using Docker on AWS EC2.
Projects:
- Microservices E-Commerce API: Built high-throughput REST APIs using Python & PostgreSQL.
- Cloud Monitoring Agent: Dockerized telemetry service deployed on AWS EC2.
Certifications: AWS Certified Cloud Practitioner
"""
(DATA_DIR / "sample_resume_arun_verma.txt").write_text(c1_txt.strip(), encoding="utf-8")

c2_txt = """
Rohan Mehta
rohan.m@example.com | +91 9777888999
Education: B.Tech Information Technology
Summary: Backend software developer with 2 years of experience in containerized backends.
Skills: Python, SQL, Docker, Git, Flask, MongoDB, Linux
Experience: 2 years in software development focusing on containerized database backends.
Projects:
- Data Ingestion Service: Flask microservice packaged with Docker Compose.
- MongoDB Aggregation Engine: SQL to NoSQL data migration pipeline.
Certifications: Docker Certified Associate
"""
(DATA_DIR / "sample_resume_rohan_mehta.txt").write_text(c2_txt.strip(), encoding="utf-8")

# Sample Job Description
jd_txt = """
Python Backend Developer
Nexora AI Solutions
Location: Remote / Hybrid

We are looking for a skilled Python Backend Developer to design, develop, and maintain high-performance microservices and REST APIs.

Key Responsibilities:
- Design robust backend architectures using Python, FastAPI, and PostgreSQL.
- Maintain source code with Git and establish automated CI/CD pipelines.
- Containerize application services using Docker and orchestrate workloads.
- Integrate cloud infrastructure on AWS (EC2, S3, RDS).

Requirements:
- Strong proficiency in Python, SQL, REST API, Git, Docker, and AWS.
- 1-2 years of relevant hands-on software development experience.
- B.Tech/B.E. in Computer Science, Information Technology, or related technical discipline.
- Knowledge of Kubernetes and Azure is a plus.
"""
(DATA_DIR / "sample_job_python_backend.txt").write_text(jd_txt.strip(), encoding="utf-8")
print("Sample documents generated successfully.")
