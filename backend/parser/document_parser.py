import io
from pathlib import Path
from typing import Union, BinaryIO, Optional, Any

def extract_text_from_pdf(file_source: Union[str, Path, bytes, BinaryIO]) -> str:
    """Extracts raw text from a PDF file using pypdf or fitz."""
    if not file_source:
        return ""

    # Ensure stream is at start if file-like
    if hasattr(file_source, "seek"):
        try:
            file_source.seek(0)
        except Exception:
            pass

    # Try pypdf first
    try:
        import pypdf
        text_content = []
        if isinstance(file_source, (str, Path)):
            reader = pypdf.PdfReader(str(file_source))
        elif isinstance(file_source, (bytes, bytearray)):
            reader = pypdf.PdfReader(io.BytesIO(file_source))
        else:
            reader = pypdf.PdfReader(file_source)
            
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_content.append(extracted)
        res = "\n".join(text_content).strip()
        if res:
            return res
    except Exception:
        pass

    # Rewind stream before attempting alternate engine
    if hasattr(file_source, "seek"):
        try:
            file_source.seek(0)
        except Exception:
            pass

    # Try PyMuPDF as alternate
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz  # type: ignore[import-not-found]
        text_content = []
        if isinstance(file_source, (str, Path)):
            doc = fitz.open(str(file_source))
        elif isinstance(file_source, (bytes, bytearray)):
            doc = fitz.open(stream=file_source, filetype="pdf")
        else:
            raw_bytes = file_source.read() if hasattr(file_source, "read") else bytes(file_source)
            doc = fitz.open(stream=raw_bytes, filetype="pdf")
            
        for page in doc:
            extracted = page.get_text()
            if extracted:
                text_content.append(extracted)
        doc.close()
        res = "\n".join(text_content).strip()
        if res:
            return res
    except Exception:
        pass

    return ""

def extract_text_from_docx(file_source: Union[str, Path, bytes, BinaryIO]) -> str:
    """Extracts raw text from a DOCX file using python-docx."""
    if not file_source:
        return ""

    if hasattr(file_source, "seek"):
        try:
            file_source.seek(0)
        except Exception:
            pass

    try:
        from docx import Document
        if isinstance(file_source, (str, Path)):
            doc = Document(str(file_source))
        elif isinstance(file_source, (bytes, bytearray)):
            doc = Document(io.BytesIO(file_source))
        else:
            doc = Document(file_source)

        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also check tables in docx with merged cell deduplication
        for table in doc.tables:
            for row in table.rows:
                seen_cells = set()
                row_texts = []
                for cell in row.cells:
                    tc_id = id(cell._tc) if hasattr(cell, '_tc') else cell
                    if tc_id not in seen_cells:
                        seen_cells.add(tc_id)
                        t = cell.text.strip() if cell.text else ""
                        if t:
                            row_texts.append(t)
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        return "\n".join(paragraphs).strip()
    except Exception:
        return ""

def extract_text(filename: Optional[str], content: Any) -> str:
    """Helper to parse either PDF, DOCX or TXT files with resilient fallback."""
    if content is None:
        return ""

    # If content is already a string and not a binary document format, return it
    if isinstance(content, str):
        if not (filename or "").lower().endswith((".pdf", ".docx", ".doc")):
            return content.strip()

    fn_lower = (filename or "").lower()
    text = ""
    
    if fn_lower.endswith(".pdf"):
        text = extract_text_from_pdf(content)
    elif fn_lower.endswith(".docx") or fn_lower.endswith(".doc"):
        text = extract_text_from_docx(content)
    elif fn_lower.endswith(".txt"):
        if isinstance(content, (bytes, bytearray)):
            text = content.decode("utf-8", errors="ignore")
        elif isinstance(content, str):
            text = content
        elif hasattr(content, "read"):
            raw = content.read()
            text = raw.decode("utf-8", errors="ignore") if isinstance(raw, bytes) else str(raw)

    # Resilient fallback if file had plain text or was simulated
    if not text.strip():
        try:
            if isinstance(content, (bytes, bytearray)):
                text = content.decode("utf-8", errors="ignore")
            elif isinstance(content, str):
                text = content
            elif hasattr(content, "read"):
                if hasattr(content, "seek"):
                    content.seek(0)
                raw = content.read()
                text = raw.decode("utf-8", errors="ignore") if isinstance(raw, bytes) else str(raw)
        except Exception:
            pass

    # Clean null characters and unprintable control characters from binary decoded text
    if '\x00' in text:
        text = "".join(ch for ch in text if ch in ('\n', '\t', '\r') or (ord(ch) >= 32 and ord(ch) != 127))

    return text.strip()
